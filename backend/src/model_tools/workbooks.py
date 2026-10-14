from __future__ import annotations

import hashlib
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

MAX_CELLS = 500_000


class WorkbookError(ValueError):
    """A workbook cannot be safely interpreted as a numeric result table."""


def column_letter(index: int) -> str:
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _scalar(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def _xlsx_extent(sheet) -> tuple[int, int]:
    """Bound actual coordinates before openpyxl allocates rectangular rows.

    Exporters can write stale/missing <dimension> metadata. Checking it alone
    can silently truncate data; blindly resetting it permits huge sparse rows.
    The XML pass uses the same source/parser as the installed openpyxl reader.
    """
    from openpyxl.utils.cell import coordinate_to_tuple
    from openpyxl.xml.functions import iterparse

    row_number = column = max_row = max_column = count = 0
    try:
        with sheet._get_source() as source:
            for event, element in iterparse(source, events=("start", "end")):
                tag = element.tag.rsplit("}", 1)[-1]
                if event == "start" and tag == "row":
                    next_row = int(element.get("r", row_number + 1))
                    if next_row <= row_number:
                        raise WorkbookError("XLSX rows must be ordered and unique")
                    row_number, column = next_row, 0
                    max_row = max(max_row, row_number)
                elif event == "start" and tag == "c":
                    r, c = coordinate_to_tuple(element.get("r")) if element.get("r") else (row_number, column + 1)
                    if r != row_number or r < 1 or c <= column:
                        raise WorkbookError("XLSX cell coordinates must match ordered rows and columns")
                    column = c
                    max_column = max(max_column, column)
                    count += 1
                if max_row * max(1, max_column) > MAX_CELLS or count > MAX_CELLS:
                    raise WorkbookError("workbook exceeds the cell limit")
                if event == "end":
                    element.clear()
    except (ValueError, KeyError) as exc:
        raise WorkbookError(f"invalid XLSX coordinates: {exc}") from exc
    return max_row, max_column


def read_workbook(path: Path, *, reject_cell_errors: bool = True) -> dict:
    """Read stored values, retaining Excel coordinates and rejecting cell errors.

    Format comes from bytes, since some supplied .xls files are actually XLSX.
    XLS formulas expose saved cached values; this reader never recalculates them.
    XLSX formulas are retained as strings and rejected by numeric validation.
    """
    if path.stat().st_size > 25 * 1024 * 1024:
        raise WorkbookError("workbook exceeds the 25 MiB local reader limit")
    raw = path.read_bytes()
    sheets = []
    total_cells = 0
    if raw.startswith(bytes.fromhex("d0cf11e0a1b11ae1")):
        import xlrd

        try:
            book = xlrd.open_workbook(file_contents=raw, on_demand=True)
        except xlrd.XLRDError as exc:
            raise WorkbookError(f"invalid XLS: {exc}") from exc
        try:
            for sheet in book.sheets():
                total_cells += sheet.nrows * sheet.ncols
                if total_cells > MAX_CELLS:
                    raise WorkbookError("workbook exceeds the cell limit")
                rows = []
                for r in range(sheet.nrows):
                    row = []
                    for c in range(sheet.ncols):
                        cell = sheet.cell(r, c)
                        if cell.ctype == xlrd.XL_CELL_ERROR and reject_cell_errors:
                            raise WorkbookError(f"Excel error at {sheet.name}!{column_letter(c+1)}{r+1}")
                        value = cell.value
                        if cell.ctype == xlrd.XL_CELL_ERROR:
                            value = xlrd.error_text_from_code.get(cell.value, "#ERROR!")
                        elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                            value = None
                        elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                            value = bool(value)
                        elif cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(value, book.datemode).isoformat()
                        row.append(value)
                    rows.append(row)
                sheets.append({"name": sheet.name, "rows": rows, "nrows": sheet.nrows, "ncols": sheet.ncols})
        finally:
            book.release_resources()
        file_format = "xls"
    elif raw.startswith(b"PK\x03\x04"):
        import openpyxl

        try:
            with ZipFile(BytesIO(raw)) as archive:
                if len(archive.infolist()) > 2000 or sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                    raise WorkbookError("XLSX exceeds the expanded archive limit")
            book = openpyxl.load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        except (ValueError, KeyError, OSError, BadZipFile, openpyxl.utils.exceptions.InvalidFileException) as exc:
            raise WorkbookError(f"invalid XLSX: {exc}") from exc
        try:
            actual_cells = 0
            for sheet in book:
                nrows, ncols = _xlsx_extent(sheet)
                total_cells += nrows * ncols
                if total_cells > MAX_CELLS:
                    raise WorkbookError("workbook exceeds the cell limit")
                rows = []
                for cells in (sheet.iter_rows(min_row=1, max_row=nrows, min_col=1, max_col=ncols) if nrows and ncols else []):
                    actual_cells += len(cells)
                    if actual_cells > MAX_CELLS:
                        raise WorkbookError("workbook exceeds the cell limit")
                    row = []
                    for cell in cells:
                        if cell.data_type == "e" and reject_cell_errors:
                            raise WorkbookError(f"Excel error at {sheet.title}!{cell.coordinate}")
                        row.append(_scalar(cell.value))
                    rows.append(row)
                sheets.append({"name": sheet.title, "rows": rows, "nrows": len(rows), "ncols": ncols})
        finally:
            book.close()
        file_format = "xlsx"
    else:
        raise WorkbookError("unsupported file content: expected XLS or XLSX")
    warnings = ["Stored workbook values only; formulas are not recalculated."]
    if path.suffix.lower() != f".{file_format}":
        warnings.append(f"Extension {path.suffix} differs from detected {file_format} content.")
    return {"file": path.name, "sha256": hashlib.sha256(raw).hexdigest(), "format": file_format,
            "warnings": warnings, "sheets": sheets}
